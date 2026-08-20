import os
import shutil
import pytest
import pymupdf
import docx
import openpyxl
from src.ingestion.ingest import process_document, ingest_all_departments

@pytest.fixture
def test_data_dir(tmp_path):
    """Creates a temporary data directory with dummy PDF, DOCX, and XLSX files."""
    data_dir = tmp_path / "data"
    
    # Create departments
    for dept in ["Finance", "HR", "IT", "Legal", "Sales"]:
        (data_dir / dept).mkdir(parents=True)
        
    # Create a dummy PDF in Finance
    pdf_path = data_dir / "Finance" / "Q1_Report.pdf"
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text((50, 50), "This is the Q1 financial report for the company. " * 50)
    doc.save(str(pdf_path))
    doc.close()

    # Create a dummy DOCX in HR
    docx_path = data_dir / "HR" / "LeavePolicy.docx"
    doc_word = docx.Document()
    doc_word.add_paragraph("This is the leave policy for Northwind Traders. " * 50)
    doc_word.save(str(docx_path))

    # Create a dummy XLSX in Sales
    xlsx_path = data_dir / "Sales" / "Pricing.xlsx"
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "2025_Prices"
    for i in range(1, 20):
        sheet.append([f"Product_{i}", i * 10, "Active"])
    wb.save(str(xlsx_path))
    
    yield str(data_dir)
    
    if os.path.exists(data_dir):
        shutil.rmtree(data_dir)

def test_process_pdf(test_data_dir):
    pdf_path = os.path.join(test_data_dir, "Finance", "Q1_Report.pdf")
    chunks = process_document(pdf_path, "Finance", chunk_size=150, chunk_overlap=20)
    
    assert len(chunks) > 0
    first_chunk = chunks[0]
    assert first_chunk["document_name"] == "Q1_Report.pdf"
    assert first_chunk["department"] == "Finance"
    assert first_chunk["page_number"] == 1
    assert first_chunk["chunk_id"].startswith("Q1_Report.pdf_Finance_p1_c")

def test_process_docx(test_data_dir):
    docx_path = os.path.join(test_data_dir, "HR", "LeavePolicy.docx")
    chunks = process_document(docx_path, "HR", chunk_size=150, chunk_overlap=20)
    
    assert len(chunks) > 0
    first_chunk = chunks[0]
    assert first_chunk["document_name"] == "LeavePolicy.docx"
    assert first_chunk["department"] == "HR"
    assert first_chunk["page_number"] == 1
    assert first_chunk["chunk_id"].startswith("LeavePolicy.docx_HR_p1_c")

def test_process_xlsx(test_data_dir):
    xlsx_path = os.path.join(test_data_dir, "Sales", "Pricing.xlsx")
    chunks = process_document(xlsx_path, "Sales")
    
    assert len(chunks) > 0
    first_chunk = chunks[0]
    assert first_chunk["document_name"] == "Pricing.xlsx"
    assert first_chunk["department"] == "Sales"
    assert first_chunk["sheet_name"] == "2025_Prices"
    assert "row_range" in first_chunk
    assert first_chunk["page_number"] == "Sheet:2025_Prices"
    # We wrote 19 rows, and chunking is every 15 rows.
    assert len(chunks) == 2

def test_ingest_all_departments(test_data_dir):
    chunks = ingest_all_departments(test_data_dir, chunk_size=200, chunk_overlap=20)
    
    assert len(chunks) > 0
    
    # Verify we got chunks from all document types
    docs = set([c["document_name"] for c in chunks])
    assert "Q1_Report.pdf" in docs
    assert "LeavePolicy.docx" in docs
    assert "Pricing.xlsx" in docs
